import os
import pandas as pd
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
from app.models import Company, QuarterlyMetrics


class ExcelHandler:
    """Handle Excel import/export operations"""

    @staticmethod
    def parse_metrics_from_excel(file_path, company_id):
        """
        Parse financial metrics from uploaded Excel file
        Expected columns: Year, Quarter, Revenue, ARR, cARR, LARR, Gross Margin, Burn Rate, Cash Balance, Customer Count
        """
        try:
            df = pd.read_excel(file_path)

            # Normalize column names (lowercase, strip whitespace)
            df.columns = df.columns.str.lower().str.strip().str.replace(' ', '_')

            metrics_data = []

            for _, row in df.iterrows():
                metric = {
                    'company_id': company_id,
                    'year': int(row.get('year', 0)) if pd.notna(row.get('year')) else None,
                    'quarter': int(row.get('quarter', 0)) if pd.notna(row.get('quarter')) else None,
                    'revenue': float(row.get('revenue', 0)) if pd.notna(row.get('revenue')) else None,
                    'arr': float(row.get('arr', 0)) if pd.notna(row.get('arr')) else None,
                    'carr': float(row.get('carr', 0)) if pd.notna(row.get('carr')) else None,
                    'larr': float(row.get('larr', 0)) if pd.notna(row.get('larr')) else None,
                    'gross_margin': float(row.get('gross_margin', 0)) if pd.notna(row.get('gross_margin')) else None,
                    'burn_rate': float(row.get('burn_rate', 0)) if pd.notna(row.get('burn_rate')) else None,
                    'cash_balance': float(row.get('cash_balance', 0)) if pd.notna(row.get('cash_balance')) else None,
                    'customer_count': int(row.get('customer_count', 0)) if pd.notna(row.get('customer_count')) else None,
                    'notes': str(row.get('notes', '')) if pd.notna(row.get('notes')) else None
                }

                # Only add if we have at least year and quarter
                if metric['year'] and metric['quarter']:
                    metrics_data.append(metric)

            return {'success': True, 'data': metrics_data, 'count': len(metrics_data)}

        except Exception as e:
            return {'success': False, 'error': str(e), 'data': []}

    @staticmethod
    def export_company_metrics(company, output_path=None):
        """
        Export a company's metrics to a formatted Excel file
        Returns the file path
        """
        from config import Config

        if output_path is None:
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            filename = f"{company.name.replace(' ', '_')}_{timestamp}.xlsx"
            output_path = os.path.join(Config.EXPORT_FOLDER, filename)

        # Fetch all metrics for the company
        metrics = QuarterlyMetrics.query.filter_by(company_id=company.id).order_by(
            QuarterlyMetrics.year, QuarterlyMetrics.quarter
        ).all()

        # Create workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Quarterly Metrics"

        # Define styles
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True, size=11)
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        # Company information header
        ws['A1'] = 'Portfolio Company Metrics Report'
        ws['A1'].font = Font(bold=True, size=14)
        ws['A2'] = f'Company: {company.name}'
        ws['A3'] = f'Industry: {company.industry or "N/A"}'
        ws['A4'] = f'Stage: {company.stage or "N/A"}'
        ws['A5'] = f'Report Generated: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}'

        # Headers starting at row 7
        headers = ['Year', 'Quarter', 'Period', 'Revenue', 'ARR', 'cARR', 'LARR',
                   'Gross Margin %', 'Burn Rate', 'Cash Balance', 'Customer Count', 'Notes']

        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=7, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = border

        # Data rows
        for row_idx, metric in enumerate(metrics, start=8):
            ws.cell(row=row_idx, column=1, value=metric.year)
            ws.cell(row=row_idx, column=2, value=metric.quarter)
            ws.cell(row=row_idx, column=3, value=f'Q{metric.quarter} {metric.year}')
            ws.cell(row=row_idx, column=4, value=metric.revenue)
            ws.cell(row=row_idx, column=5, value=metric.arr)
            ws.cell(row=row_idx, column=6, value=metric.carr)
            ws.cell(row=row_idx, column=7, value=metric.larr)
            ws.cell(row=row_idx, column=8, value=metric.gross_margin)
            ws.cell(row=row_idx, column=9, value=metric.burn_rate)
            ws.cell(row=row_idx, column=10, value=metric.cash_balance)
            ws.cell(row=row_idx, column=11, value=metric.customer_count)
            ws.cell(row=row_idx, column=12, value=metric.notes or '')

            # Apply borders and number formatting
            for col in range(1, 13):
                cell = ws.cell(row=row_idx, column=col)
                cell.border = border

                # Format currency columns
                if col in [4, 5, 6, 7, 9, 10]:
                    cell.number_format = '$#,##0.00'
                # Format percentage
                elif col == 8:
                    cell.number_format = '0.00%'

        # Adjust column widths
        column_widths = [8, 8, 12, 15, 15, 15, 15, 15, 15, 15, 15, 30]
        for col, width in enumerate(column_widths, start=1):
            ws.column_dimensions[chr(64 + col)].width = width

        # Save workbook
        wb.save(output_path)
        return output_path

    @staticmethod
    def export_all_companies(output_path=None):
        """
        Export all companies and their metrics to a comprehensive Excel file
        """
        from config import Config

        if output_path is None:
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            filename = f"portfolio_overview_{timestamp}.xlsx"
            output_path = os.path.join(Config.EXPORT_FOLDER, filename)

        companies = Company.query.all()
        wb = Workbook()

        # Remove default sheet
        wb.remove(wb.active)

        # Overview sheet
        ws_overview = wb.create_sheet("Portfolio Overview")
        overview_headers = ['Company Name', 'Industry', 'Stage', 'Investment Date',
                           'Latest Revenue', 'Latest ARR', 'Latest Quarter']

        for col, header in enumerate(overview_headers, start=1):
            ws_overview.cell(row=1, column=col, value=header).font = Font(bold=True)

        for row_idx, company in enumerate(companies, start=2):
            latest_metric = QuarterlyMetrics.query.filter_by(company_id=company.id).order_by(
                QuarterlyMetrics.year.desc(), QuarterlyMetrics.quarter.desc()
            ).first()

            ws_overview.cell(row=row_idx, column=1, value=company.name)
            ws_overview.cell(row=row_idx, column=2, value=company.industry)
            ws_overview.cell(row=row_idx, column=3, value=company.stage)
            ws_overview.cell(row=row_idx, column=4, value=company.investment_date)

            if latest_metric:
                ws_overview.cell(row=row_idx, column=5, value=latest_metric.revenue)
                ws_overview.cell(row=row_idx, column=6, value=latest_metric.arr)
                ws_overview.cell(row=row_idx, column=7, value=f'Q{latest_metric.quarter} {latest_metric.year}')

        # Individual company sheets
        for company in companies:
            metrics = QuarterlyMetrics.query.filter_by(company_id=company.id).order_by(
                QuarterlyMetrics.year, QuarterlyMetrics.quarter
            ).all()

            if metrics:
                ws = wb.create_sheet(company.name[:31])  # Excel sheet name limit is 31 chars

                headers = ['Period', 'Revenue', 'ARR', 'cARR', 'LARR',
                          'Gross Margin', 'Burn Rate', 'Cash Balance', 'Customers']

                for col, header in enumerate(headers, start=1):
                    ws.cell(row=1, column=col, value=header).font = Font(bold=True)

                for row_idx, metric in enumerate(metrics, start=2):
                    ws.cell(row=row_idx, column=1, value=f'Q{metric.quarter} {metric.year}')
                    ws.cell(row=row_idx, column=2, value=metric.revenue)
                    ws.cell(row=row_idx, column=3, value=metric.arr)
                    ws.cell(row=row_idx, column=4, value=metric.carr)
                    ws.cell(row=row_idx, column=5, value=metric.larr)
                    ws.cell(row=row_idx, column=6, value=metric.gross_margin)
                    ws.cell(row=row_idx, column=7, value=metric.burn_rate)
                    ws.cell(row=row_idx, column=8, value=metric.cash_balance)
                    ws.cell(row=row_idx, column=9, value=metric.customer_count)

        wb.save(output_path)
        return output_path
